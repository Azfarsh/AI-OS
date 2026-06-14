#!/usr/bin/env python3
"""
generate_contract.py — Bombay Media Contract Generator
Clones the branded PPTX template and fills all placeholders.
Outputs a filled PPTX + PDF to clients/{slug}/contracts/

Usage:
    python scripts/generate_contract.py \
        --client-name "Acme Coaching" \
        --client-email "john@acmecorp.com" \
        --client-city "New York" \
        --client-country "USA" \
        --client-contact "John Smith" \
        --client-title "CEO" \
        --client-website "acmecorp.com" \
        --services "meta,content" \
        --monthly-retainer "3000" \
        --ad-spend-min "3000" \
        --ad-spend-max "5000" \
        --agreement-number "BM-2025-002"
"""

import argparse
import os
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

# ── python-pptx ────────────────────────────────────────────────────
try:
    from pptx import Presentation
    from pptx.util import Pt
except ImportError:
    print("ERROR: python-pptx not installed. Run: pip install python-pptx")
    sys.exit(1)


# ── Path config ────────────────────────────────────────────────────
REPO_ROOT = Path(__file__).parent.parent
TEMPLATE_PPTX = REPO_ROOT / "references" / "assets" / "Bombay_Media_Contract_Template.pptx"
CLIENTS_DIR = REPO_ROOT / "clients"


def slugify(name: str) -> str:
    return name.lower().strip().replace(" ", "-").replace("'", "").replace(",", "")


def services_label(services_str: str) -> str:
    """Convert 'meta,content' → 'Meta Ads Management and Content Marketing'"""
    mapping = {
        "meta": "Meta Ads Management",
        "google": "Google Ads Management",
        "content": "Content Marketing",
        "seo": "SEO & Content",
    }
    parts = [mapping.get(s.strip(), s.strip().title()) for s in services_str.split(",")]
    if len(parts) == 1:
        return parts[0]
    return ", ".join(parts[:-1]) + " and " + parts[-1]


def replace_text_in_shape(shape, replacements: dict):
    """Replace placeholder tokens in a shape's text frame, preserving formatting."""
    if not shape.has_text_frame:
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            for placeholder, value in replacements.items():
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, value)


def fill_contract(args) -> Path:
    slug = slugify(args.client_name)
    date_str = datetime.today().strftime("%d %B %Y")
    year = datetime.today().strftime("%Y")

    # Auto-generate agreement number if not provided
    agreement_no = args.agreement_number or f"BM-{year}-{datetime.today().strftime('%m%d')}"

    # Build replacements dict
    services_text = services_label(args.services)
    ad_spend_range = f"${int(args.ad_spend_min):,}–${int(args.ad_spend_max):,}/month"
    monthly_retainer = f"${int(args.monthly_retainer):,}/month"

    replacements = {
        "[CLIENT COMPANY NAME]": args.client_name,
        "[DD MONTH YYYY]": date_str,
        "BM-2025-001": agreement_no,
        "[Client Company Name]": args.client_name,
        "[City, Country]": f"{args.client_city}, {args.client_country}",
        "[Primary Contact Name]": args.client_contact,
        "[Title]": args.client_title,
        "[Email Address]": args.client_email,
        "[Website URL]": args.client_website,
        "[Name], [Title]": f"{args.client_contact}, {args.client_title}",
        "[Client Company]": args.client_name,
        # Fee placeholders (slide 3 table)
        "$3,000/month": monthly_retainer,
        "$3,000–$10,000/month": ad_spend_range,
    }

    # ── Load template ──────────────────────────────────────────────
    if not TEMPLATE_PPTX.exists():
        print(f"ERROR: Template not found at {TEMPLATE_PPTX}")
        print("  → Copy Bombay_Media_Dummy_Contract.pptx to references/assets/Bombay_Media_Contract_Template.pptx")
        sys.exit(1)

    prs = Presentation(str(TEMPLATE_PPTX))

    # ── Fill all slides ────────────────────────────────────────────
    for slide in prs.slides:
        for shape in slide.shapes:
            replace_text_in_shape(shape, replacements)
            # Handle table cells (slide 3 fee table)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for placeholder, value in replacements.items():
                            if placeholder in cell.text_frame.text:
                                for para in cell.text_frame.paragraphs:
                                    for run in para.runs:
                                        if placeholder in run.text:
                                            run.text = run.text.replace(placeholder, value)

    # ── Output paths ───────────────────────────────────────────────
    contract_dir = CLIENTS_DIR / slug / "contracts"
    contract_dir.mkdir(parents=True, exist_ok=True)

    date_tag = datetime.today().strftime("%Y-%m-%d")
    pptx_path = contract_dir / f"contract-{date_tag}.pptx"
    pdf_path = contract_dir / f"contract-{date_tag}.pdf"

    prs.save(str(pptx_path))
    print(f"OK Contract PPTX saved: {pptx_path}")

    # ── Convert to PDF (requires LibreOffice) ──────────────────────
    try:
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf",
             "--outdir", str(contract_dir), str(pptx_path)],
            capture_output=True, text=True, timeout=60
        )
        if result.returncode == 0:
            # soffice names the PDF after the pptx filename
            generated_pdf = contract_dir / f"contract-{date_tag}.pdf"
            print(f"OK Contract PDF saved: {generated_pdf}")
            return generated_pdf
        else:
            print(f"WARN PDF conversion failed: {result.stderr}")
            print(f"  -> PPTX available at: {pptx_path}")
            return pptx_path
    except (FileNotFoundError, subprocess.TimeoutExpired):
        print("WARN LibreOffice not found - PDF conversion skipped.")
        print(f"  -> PPTX available at: {pptx_path}")
        return pptx_path


def main():
    parser = argparse.ArgumentParser(description="Bombay Media — Contract Generator")
    parser.add_argument("--client-name",       required=True,  help="e.g. 'Acme Coaching'")
    parser.add_argument("--client-email",      required=True,  help="e.g. 'john@acme.com'")
    parser.add_argument("--client-city",       required=True,  help="e.g. 'New York'")
    parser.add_argument("--client-country",    required=True,  help="e.g. 'USA'")
    parser.add_argument("--client-contact",    required=True,  help="e.g. 'John Smith'")
    parser.add_argument("--client-title",      default="Founder", help="e.g. 'CEO'")
    parser.add_argument("--client-website",    default="",     help="e.g. 'acme.com'")
    parser.add_argument("--services",          required=True,  help="Comma-separated: meta,content,google,seo")
    parser.add_argument("--monthly-retainer",  default="3000", help="Monthly retainer in USD (number only)")
    parser.add_argument("--ad-spend-min",      default="3000", help="Min ad spend in USD (number only)")
    parser.add_argument("--ad-spend-max",      default="10000",help="Max ad spend in USD (number only)")
    parser.add_argument("--agreement-number",  default=None,   help="e.g. BM-2025-002 (auto-generated if blank)")

    args = parser.parse_args()
    output_path = fill_contract(args)
    # Print the path for the calling skill to capture
    print(f"CONTRACT_PATH:{output_path}")


if __name__ == "__main__":
    main()
